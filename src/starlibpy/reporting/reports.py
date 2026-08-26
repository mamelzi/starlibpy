"""Multi-result report assembly."""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from starlibpy.results import StarFigure, StarTable

from .exporters import export_plot, export_table
from .plots import plot_result
from .tables import render_table


@dataclass
class StarReport:
    title: str
    sections: list[dict[str, Any]] = field(default_factory=list)

    def add_result(
        self,
        result: Any,
        *,
        title: str | None = None,
        table: str | None = None,
        plot: str | None = None,
    ):
        section = {"title": title, "result": result, "table": render_table(result, name=table)}
        if plot is not False:
            try:
                section["figure"] = plot_result(result, kind=plot)
            except Exception:
                pass
        self.sections.append(section)
        return self


def render_report(results: Sequence[Any], *, title: str = "Starlibpy report") -> StarReport:
    report = StarReport(title)
    for result in results:
        report.add_result(result)
    return report


def export_report(report: StarReport, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".html":
        parts = [f"<h1>{report.title}</h1>"]
        for i, s in enumerate(report.sections):
            parts.append(f"<h2>{s.get('title') or f'Section {i + 1}'}</h2>")
            parts.append(s["table"].data.to_html(index=False))
        path.write_text("\n".join(parts), encoding="utf-8")
        return path
    if path.suffix.lower() == ".docx":
        from docx import Document

        doc = Document()
        doc.add_heading(report.title, 0)
        for i, s in enumerate(report.sections):
            doc.add_heading(s.get("title") or f"Section {i + 1}", level=1)
            data = s["table"].data
            t = doc.add_table(rows=1, cols=len(data.columns))
            t.style = "Table Grid"
            for j, c in enumerate(data.columns):
                t.rows[0].cells[j].text = str(c)
            for _, row in data.iterrows():
                cells = t.add_row().cells
                [setattr(cells[j], "text", str(v)) for j, v in enumerate(row)]
        doc.save(path)
        return path
    raise ValueError("Report export currently supports HTML and DOCX.")
