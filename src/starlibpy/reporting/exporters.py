"""Table, figure, result, and report exporters."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd

from starlibpy.results import StarFigure, StarTable


def export_table(
    table: StarTable | pd.DataFrame, path: str | Path, *, index: bool = False, **kwargs
) -> Path:
    """Export a rendered table based on the destination extension."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = table.data if isinstance(table, StarTable) else table
    ext = path.suffix.lower()
    if ext == ".csv":
        data.to_csv(path, index=index, **kwargs)
    elif ext in {".xlsx", ".xlsm"}:
        data.to_excel(path, index=index, engine="openpyxl", **kwargs)
    elif ext in {".md", ".markdown"}:
        path.write_text(data.to_markdown(index=index), encoding="utf-8")
    elif ext in {".tex", ".latex"}:
        path.write_text(data.to_latex(index=index), encoding="utf-8")
    elif ext in {".html", ".htm"}:
        path.write_text(data.to_html(index=index, border=0), encoding="utf-8")
    elif ext == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ImportError("Word export requires python-docx.") from exc
        doc = Document()
        if isinstance(table, StarTable) and table.title:
            doc.add_heading(table.title, level=1)
        t = doc.add_table(rows=1, cols=len(data.columns))
        t.style = "Table Grid"
        for i, c in enumerate(data.columns):
            t.rows[0].cells[i].text = str(c)
        for _ in range(len(data)):
            t.add_row()
        # populate in a second clear loop to avoid stale row reference
        for ri, (_, row) in enumerate(data.iterrows(), start=1):
            for ci, v in enumerate(row):
                t.rows[ri].cells[ci].text = "" if pd.isna(v) else str(v)
        if isinstance(table, StarTable):
            for note in table.notes:
                doc.add_paragraph(str(note), style=None)
        doc.save(path)
    else:
        raise ValueError("Supported table extensions: csv, xlsx, md, tex, html, docx.")
    return path


def export_plot(
    figure: StarFigure | Any,
    path: str | Path,
    *,
    dpi: int = 300,
    transparent: bool = False,
    **kwargs,
) -> Path:
    """Export a figure to a Matplotlib-supported image format."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fig = figure.figure if isinstance(figure, StarFigure) else figure
    fig.savefig(path, dpi=dpi, transparent=transparent, bbox_inches="tight", **kwargs)
    return path


def export_result(
    result: Any,
    path: str | Path,
    *,
    template: str = "journal",
    theme: str = "default",
    include_plots: bool = True,
) -> Path:
    """Export all declared tables and optional plots to a directory."""
    from .plots import plot_result
    from .tables import render_table

    path = Path(path)
    path.mkdir(parents=True, exist_ok=True)
    for name in getattr(result, "available_tables", ()):
        try:
            export_table(
                render_table(result, name=name, template=template, theme=theme),
                path / f"{name}.xlsx",
            )
        except Exception:
            result.get_table(name).to_csv(path / f"{name}.csv", index=False)
    if include_plots:
        for kind in getattr(result, "available_plots", ()):
            try:
                export_plot(plot_result(result, kind=kind, theme=theme), path / f"{kind}.svg")
            except Exception:
                pass
    metadata = (
        result.to_dict(include_models=False)
        if hasattr(result, "to_dict")
        else {"repr": repr(result)}
    )
    (path / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    return path


def export_analysis_bundle(result: Any, path: str | Path, **kwargs) -> Path:
    return export_result(result, path, **kwargs)
